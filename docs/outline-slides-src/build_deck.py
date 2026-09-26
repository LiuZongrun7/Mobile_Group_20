#!/usr/bin/env python3
"""Build the TokenTrail outline presentation (.pptx).

Content is taken from TokenTrail_Project_Outline.tex so the deck and the
2-4 page summary tell the same story; the six numbered sections mirror the
six bullet points the outline guidelines ask the slides to cover.
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

# ---------------------------------------------------------------- palette
NAVY, TEAL, AMBER = '315E75', '2E766E', 'A17431'
INK, MUTE = '2B3A42', '7A8C94'
LIGHTBLUE, LIGHTTEAL, LIGHTAMBER = 'EAF3F7', 'EDF6F2', 'FAF3E5'
LIGHTROW, HAIRLINE, WHITE = 'F7FAFB', 'BDD0D8', 'FFFFFF'
FONT = 'Arial'
GH = 'https://github.com/LiuZongrun7/Mobile_Group_20'

BASE = os.path.dirname(os.path.abspath(__file__))
MARGIN, PAGE_W, PAGE_H = 0.72, 13.333, 7.5
CW = PAGE_W - 2 * MARGIN          # content width = 11.893
FOOT_Y = 7.02

_prs = Presentation()
_prs.slide_width, _prs.slide_height = Inches(PAGE_W), Inches(PAGE_H)
BLANK = _prs.slide_layouts[6]
_slide_no = [0]


def rgb(h):
    return RGBColor.from_string(h)


def tb(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, first, parts, size=13, color=INK, align=PP_ALIGN.LEFT,
         before=0, after=5, line=1.16, bold=False, italic=False):
    """parts: str, or list of (text, bold, color, italic, href) tuples."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before, p.space_after = Pt(before), Pt(after)
    p.line_spacing = line
    if isinstance(parts, str):
        parts = [(parts, bold, color, italic, None)]
    for part in parts:
        text, bold, col, italic, href = (list(part) + [None] * 5)[:5]
        r = p.add_run()
        r.text = text
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), bool(bold), bool(italic)
        f.color.rgb = rgb(col or color)
        if href:
            r.hyperlink.address = href
            f.underline = False
    return p


def shape(slide, kind, l, t, w, h, fill=None, line=None, lw=0.75, radius=0.08):
    s = slide.shapes.add_shape(kind, Inches(l), Inches(t), Inches(w), Inches(h))
    s.shadow.inherit = False
    # Drop the theme shape style too: some renderers (LibreOffice, older WPS)
    # still apply its effectRef and draw a drop shadow behind every card.
    style = s._element.find(qn('p:style'))
    if style is not None:
        s._element.remove(style)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.word_wrap = True
    return s


def rule(slide, l, t, w, color=HAIRLINE, h=0.014):
    return shape(slide, MSO_SHAPE.RECTANGLE, l, t, w, h, fill=color)


def slide(eyebrow=None, title=None, num=True):
    s = _prs.slides.add_slide(BLANK)
    _slide_no[0] += 1
    if num:
        rule(s, MARGIN, 6.90, CW, HAIRLINE)
        tf = tb(s, MARGIN, FOOT_Y, CW * 0.7, 0.3)
        para(tf, True, 'TokenTrail · Group 20 · Android App Project Outline',
             size=9, color=MUTE, after=0)
        tf = tb(s, PAGE_W - MARGIN - 1.2, FOOT_Y, 1.2, 0.3)
        para(tf, True, str(_slide_no[0]), size=9, color=MUTE,
             align=PP_ALIGN.RIGHT, after=0)
    if eyebrow:
        tf = tb(s, MARGIN, 0.40, CW, 0.26)
        para(tf, True, eyebrow.upper(), size=10, color=AMBER, after=0)
    if title:
        tf = tb(s, MARGIN, 0.66, CW, 0.55)
        para(tf, True, title, size=26, color=NAVY, after=0)
        rule(s, MARGIN, 1.34, CW, NAVY, h=0.025)
    return s


def card(slide, l, t, w, h, head=None, body=None, fill=LIGHTROW,
         line=HAIRLINE, head_color=NAVY, size=11.5, head_size=12.5,
         tag=None, tag_color=AMBER, pad=0.17, foot=None, foot_color=MUTE):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h,
              fill=fill, line=line, radius=0.06)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(pad * 0.8)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    first = True
    if tag:
        para(tf, True, tag.upper(), size=8.5, color=tag_color, after=3)
        first = False
    if head:
        para(tf, first, head, size=head_size, bold=True, color=head_color, after=3)
        first = False
    if body:
        lines = body if isinstance(body, list) else [body]
        for i, text in enumerate(lines):
            para(tf, first, text, size=size, after=0 if i == len(lines) - 1 else 5,
                 line=1.14)
            first = False
    if foot:
        para(tf, first, foot, size=size - 1.5, color=foot_color, before=6,
             after=0, line=1.12)
    return s


# ------------------------------------------------------------------ tables
CT_ORDER = ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB', 'a:lnTlToBr', 'a:lnBlToTr',
            'a:cell3D', 'a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill',
            'a:pattFill', 'a:grpFill', 'a:headers', 'a:extLst']
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NO_STYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'


def _ordered_insert(tcPr, el, tag):
    idx = CT_ORDER.index(tag)
    for child in tcPr:
        ctag = 'a:' + child.tag.split('}')[-1]
        if ctag in CT_ORDER and CT_ORDER.index(ctag) > idx:
            child.addprevious(el)
            return
    tcPr.append(el)


def cell_border(cell, color=HAIRLINE, width=0.75, edges='LRTB'):
    tcPr = cell._tc.get_or_add_tcPr()
    for e in edges:
        tag = 'a:ln' + e
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
        _ordered_insert(tcPr, parse_xml(
            '<%s xmlns:a="%s" w="%d" cap="flat" cmpd="sng" algn="ctr">'
            '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
            '<a:prstDash val="solid"/></%s>'
            % (tag, A_NS, int(Pt(width)), color, tag)), tag)


def table(slide, l, t, col_w, rows, row_h=0.42, size=10.5, head_size=11,
          head_fill=NAVY, zebra=True, head_h=0.44):
    nrows, ncols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nrows, ncols, Inches(l), Inches(t),
                                Inches(sum(col_w)),
                                Inches(head_h + row_h * (nrows - 1)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    st = tblPr.find(qn('a:tableStyleId'))
    if st is None:
        st = parse_xml('<a:tableStyleId xmlns:a="%s"/>' % A_NS)
        tblPr.append(st)
    st.text = NO_STYLE
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Inches(w)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(head_h if ri == 0 else row_h)
        for ci, val in enumerate(row):
            opts = {}
            if isinstance(val, tuple):
                val, opts['href'] = val
            elif isinstance(val, list):          # several runs, each optionally linked
                opts['runs'] = val
            c = tbl.cell(ri, ci)
            c.margin_left, c.margin_right = Inches(0.10), Inches(0.09)
            c.margin_top = c.margin_bottom = Inches(0.055)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(
                head_fill if ri == 0 else
                (WHITE if (ri % 2 or not zebra) else LIGHTROW))
            cell_border(c, color=head_fill if ri == 0 else HAIRLINE,
                        width=0.75, edges='TB')
            if ci == 0:
                cell_border(c, color=head_fill if ri == 0 else HAIRLINE,
                            edges='L')
            if ci == ncols - 1:
                cell_border(c, color=head_fill if ri == 0 else HAIRLINE,
                            edges='R')
            tf = c.text_frame
            tf.word_wrap = True
            is_head = ri == 0
            href = opts.get('href')
            col = WHITE if is_head else (
                NAVY if (ci == 0 and ncols > 2 and not is_head) else INK)
            if 'runs' in opts:
                parts = [(t, False, (c[0] if c else INK), False, h)
                         for t, h, *c in [tuple(r) + (None,) * (3 - len(r))
                                          for r in opts['runs']]]
            else:
                parts = [(str(val), True, col, False, href)] if (is_head or ci == 0 or href) \
                    else [(str(val), False, col, False, None)]
            para(tf, True, parts, size=head_size if is_head else size,
                 color=col, after=0, line=1.08)
    return tbl


# ================================================================ 1. title
s = slide(num=False)
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, PAGE_W, 0.16, fill=NAVY)
tf = tb(s, MARGIN, 1.55, 8.6, 0.3)
para(tf, True, 'MOBILE APP PROJECT · OUTLINE PRESENTATION', size=11, color=AMBER, after=0)
tf = tb(s, MARGIN, 1.95, 8.6, 1.1)
para(tf, True, 'TokenTrail', size=54, bold=True, color=NAVY, after=0)
tf = tb(s, MARGIN, 3.10, 8.6, 0.7)
para(tf, True, 'Track, price and compare what your coding agents really cost',
     size=19, color=TEAL, after=0)
rule(s, MARGIN, 3.85, 8.6, HAIRLINE)
tf = tb(s, MARGIN, 4.05, 8.6, 1.5)
para(tf, True, [('Category:  ', True, NAVY), ('Developer Productivity / API Cost Analytics', False, INK)],
     size=13, after=6)
para(tf, False, [('Target users:  ', True, NAVY),
                 ('Students and independent developers who use API-key coding agents', False, INK)],
     size=13, after=6)
para(tf, False, [('Stack:  ', True, NAVY),
                 ('Android · Java/XML · Room · Java advice service on the team\'s own server',
                  False, INK)], size=13, after=0)
tf = tb(s, MARGIN, 5.75, 8.6, 0.9)
para(tf, True, [('Group 20 ·  ', True, NAVY),
                ('Zhang Li (24107757) · Wang Tingdong (24107759) · Liu Zongrun (24107745)',
                 False, INK)], size=12.5, after=4)
para(tf, False, [('Repository:  ', True, NAVY), (GH.replace('https://', ''), False, TEAL, False, GH)],
     size=12.5, after=0)
s.shapes.add_picture(os.path.join(BASE, 'fig_game.png'), Inches(9.85), Inches(1.55), height=Inches(4.95))

# ================================================= 2. what the slides cover
s = slide('How to read this deck', 'The six points the outline slides must cover')
items = [
    ('1', 'Brief app idea description', 'What TokenTrail is, who it is for, and what is out of scope', 'Slide 3'),
    ('2', 'Discussion of similar / related apps', 'Codex, ZCode, DSH with dsh-context, LiteLLM and dsh-pet', 'Slide 5'),
    ('3', 'Related apps: features and components', 'What each one covers, and the gap it leaves open', 'Slide 6'),
    ('4', 'Open-source code available to reuse', 'Libraries, samples and licences we will record', 'Slide 7'),
    ('5', 'Proposed features and uniqueness', 'Analytics and the home-screen assistant as the core; forum and game as supporting tabs', 'Slides 8-13'),
    ('6', 'Approach and work plan to week 15', 'Alpha, Beta and Final checkpoints with acceptance tests', 'Slides 14-15'),
]
table(s, MARGIN, 1.62, [3.55, 6.45, 1.893],
      [['Guideline point', 'What the slides cover', 'Where']]
      + [['%s.  %s' % (n, head), body, where] for n, head, body, where in items],
      row_h=0.70, head_h=0.46, size=11.5, head_size=11.5)
tf = tb(s, MARGIN, 6.42, CW, 0.4)
para(tf, True, 'The written outline covers the same six points in the same order; this deck is the visual version of it.',
     size=10.5, color=MUTE, after=0)

# ============================================================ 3. app idea
s = slide('1 · App idea', 'An API-cost analytics app that makes agent spend visible')
tf = tb(s, MARGIN, 1.62, CW, 0.75)
para(tf, True, [('TokenTrail is an Android API-cost app for people who use several coding agents: '
                 'Codex with OpenAI, ZCode with Xiaomi MiMo and DSH with DeepSeek. Its analytics home screen '
                 'prices the pay-as-you-go calls those agents already logged, and an assistant on the same '
                 'screen answers cost questions with evidence.', False, INK)],
     size=14, after=0, line=1.2)
parts = [
    ('API-cost analytics (home)', 'Prices logged pay-as-you-go calls from Codex, ZCode and DSH at dated official rates, with source, coverage and session detail wherever the logs support it.', LIGHTTEAL, TEAL),
    ('Home-screen AI assistant', 'An "Ask AI / hold to talk" entry above the navigation. Type, or dictate and confirm the transcript; replies cite evidence and admit missing data. Its own API cost is kept separate.', LIGHTBLUE, NAVY),
    ('Supporting tabs', 'Forum: official model news plus user posts and tips. Game: a monthly tower-defence season built from tokens already logged. Both are extensions, not the main experience.', LIGHTAMBER, AMBER),
]
cw, gap = (CW - 2 * 0.26) / 3, 0.26
for i, (head, body, fill, hc) in enumerate(parts):
    card(s, MARGIN + i * (cw + gap), 2.46, cw, 1.92, head=head, body=body,
         fill=fill, head_color=hc, size=11, head_size=13)
agent = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, 4.58, CW, 1.00, fill=NAVY, radius=0.06)
tfa = agent.text_frame
tfa.margin_left = tfa.margin_right = Inches(0.22)
tfa.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tfa, True, [('Scope rule.  ', True, WHITE),
                 ('API-cost analytics and the home-screen assistant are core; forum and game are supporting '
                  'extensions. Token Plan and other subscriptions are excluded, and TokenTrail never runs, '
                  'routes or controls a coding agent.', False, 'D6E4EB')], size=13, after=0, line=1.15)
card(s, MARGIN, 5.78, CW, 0.92,
     head='Estimated, never invoiced',
     body='An API key alone cannot recover an agent\'s history, so a call is priced only when a saved rate snapshot '
          'covers its date. Missing logs or rates stay visible, and every monetary value is an estimate.',
     fill=LIGHTROW, head_color=NAVY, size=11, head_size=12)

# ============================================================== 4. problem
s = slide('1 · App idea', 'Why this is worth building')
probs = [
    ('Each console sees one account', 'Codex, ZCode and DSH bill separately, in their own units, with no cross-provider total and no common price version. A user running three agents has three bills and no single monthly figure.'),
    ('Tokens are not money', 'A response reports tokens for one call, not the agent\'s whole history. A call can only be priced when a saved rate snapshot covers its date, so some calls stay Unavailable rather than being guessed at.'),
    ('Rating your own work is bad input', 'Asking users to score a task costs them extra effort and produces opinions, not measurements. It also makes the numbers depend on mood instead of on the data.'),
    ('Spend leaves no trace', 'A day of agent work leaves nothing visible behind, so usage never turns into progress the user can see and nothing brings them back to the tools.'),
]
cw, ch, gap = (CW - 0.30) / 2, 1.80, 0.30
for i, (head, body) in enumerate(probs):
    card(s, MARGIN + (i % 2) * (cw + gap), 1.68 + (i // 2) * (ch + 0.22), cw, ch,
         head=head, body=body, fill=LIGHTROW, size=11.5, head_size=13.5)
card(s, MARGIN, 5.72, CW, 1.04, head='What TokenTrail does instead',
     body='It computes only what the data supports — logged token counts, dated official rates, cache behaviour and '
          'duration — and labels every figure Estimated, Partial estimate or Unavailable. It never asks the user to '
          'judge how good an answer was.',
     fill=LIGHTBLUE, head_color=NAVY, size=11.5, head_size=12.5)

# ========================================================= 5. related apps
s = slide('2 · Related apps', 'Five products we studied before writing a line of code')
apps = [
    ('FIRST-PARTY AGENT', 'Codex', 'OpenAI\'s coding agent. Can run with an OpenAI API key and be billed through the Platform account.'),
    ('MODEL PROVIDER', 'ZCode', 'Supports Xiaomi MiMo as a model provider. Usable MiMo API logs are imported; Token Plan usage is excluded.'),
    ('SESSION PLUGIN', 'DSH / dsh-context', 'DSH runs coding sessions; its dsh-context plugin visualises tokens, cache and activity.'),
    ('GATEWAY', 'LiteLLM', 'Open-source gateway with spend tracking and budgets — but only for traffic that goes through it.'),
    ('COMPANION', 'dsh-pet', 'Floating companion driven by DSH session events, with touch and feeding actions. Closest thing to a game.'),
]
feet = ['The OpenAI-side agent we track.',
        'Tracked through its MiMo API endpoint.',
        'Import its logs into Android; no live agent control.',
        'Teaches the vocabulary of spend tracking and budgets.',
        'Its entry point inspires our assistant panel on the home screen.']
gap = 0.16
cw = (CW - 4 * gap) / 5
for i, (tag, name, body) in enumerate(apps):
    card(s, MARGIN + i * (cw + gap), 1.74, cw, 2.42, head=name, body=body,
         tag=tag, size=11, head_size=15 if len(name) < 10 else 13,
         fill=LIGHTROW, pad=0.15, foot=feet[i])
left, rw, rgap = CW * 0.60, CW * 0.40 - 0.24, 0.24
card(s, MARGIN, 4.36, left, 2.06, head='What the review told us',
     body=['Every product above is strong inside its own boundary: a first-party console prices one provider, a plugin '
           'reads one harness, a gateway sees only proxied traffic. None of them prices three providers from records '
           'the user already owns, and only dsh-pet treats usage as something to come back to.',
           'The review also set our boundaries: no proxying, no subscription quotas, and no overlay on another app.'],
     fill=LIGHTTEAL, head_color=TEAL, size=12, head_size=13)
card(s, MARGIN + left + rgap, 4.36, rw, 2.06, head='The opportunity',
     body=['The outline names three contributions:',
           '▪  unified analytics over logs the user owns;',
           '▪  reproducible costs that keep the price version;',
           '▪  evidence-based advice that admits missing data.'],
     fill=LIGHTAMBER, head_color=AMBER, size=11.5, head_size=13)
tf = tb(s, MARGIN, 6.54, CW, 0.3)
para(tf, True, 'Section 2 of the written outline compares the same five products, with a link to each one.',
     size=10.5, color=MUTE, after=0)

# ================================================= 6. related apps: table
s = slide('3 · Related apps in detail', 'Component by component: what they cover, what is left open')
table(s, MARGIN, 1.62,
      [1.75, 5.05, 5.093],
      [['Product', 'Relevant capability', 'Gap addressed by TokenTrail'],
       [('Codex', 'https://learn.chatgpt.com/docs/auth'),
        'API-key coding agent billed through OpenAI Platform.',
        'Dated estimates from the available Codex logs.'],
       [('ZCode', 'https://zcode.z.ai/en/docs/configuration'),
        'Supports Xiaomi MiMo as a model provider.',
        'Import usable MiMo API logs; exclude Token Plan usage.'],
       [('DSH / dsh-context', 'https://github.com/deepseek-ai/deepseek-harness'),
        'DSH runs coding sessions; its context plugin visualises tokens, cache and activity.',
        'Import usable DSH logs into Android; no live agent control.'],
       [('LiteLLM', 'https://docs.litellm.ai/docs/simple_proxy'),
        'Gateway spend tracking and budgets for proxied calls.',
        'Imports records without proxying agent traffic.'],
       [('dsh-pet', 'https://github.com/zhu1090093659/dsh-pet/blob/main/README.zh.md'),
        'Floating companion driven by DSH session events.',
        'Interactive home-screen assistant with questions and evidence-linked advice.']],
      row_h=0.80, head_h=0.46, size=10.5, head_size=11.5)
tf = tb(s, MARGIN, 6.20, CW, 0.45)
para(tf, True, 'Names link to each product\'s own documentation or repository. The full comparison is section 2 of the written outline.',
     size=10.5, color=MUTE, after=0)

# =========================================================== 7. open source
s = slide('4 · Open source and tools', 'What we reuse, and what we build ourselves')
table(s, MARGIN, 1.62, [2.45, 9.443],
      [['Layer', 'Planned tools and rationale'],
       ['Android interface',
        'Java/XML, Material Components, Navigation and ViewModel/LiveData; RecyclerView lists, '
        'speech-to-text and typed fallback.'],
       ['Network and work',
        'Retrofit/OkHttp connects the Java advice service; bounded read-only tools and one model API. '
        'WorkManager checks prices.'],
       ['Persistence and security',
        'Room stores usage, prices and history; the team\'s own server supports sync and public posts. '
        'The service authenticates every request and scopes every read to the caller\'s account; '
        'no client holds a model key.'],
       ['Open-source reuse',
        [('Room', 'https://github.com/androidx/androidx/blob/androidx-main/LICENSE.txt', TEAL),
         (' (Apache-2.0) for local records; ', None),
         ('MPAndroidChart', 'https://github.com/PhilJay/MPAndroidChart', TEAL),
         (' (Apache-2.0) for charts, added as a dependency; ', None),
         ('Room with a View', 'https://github.com/android/codelab-android-room-with-a-view', TEAL),
         (' is an archived architecture reference with no copied code. dsh-context/LiteLLM: comparison only; '
          'dsh-pet: entry inspiration only. The game engine is original.', None)]]],
      row_h=0.74, head_h=0.46, size=10.5)
card(s, MARGIN, 5.16, CW, 0.78,
     body='Original or properly licensed art only. Every reused file and its licence, version and source is '
          'recorded in the repository, and this layer is revisited if a dependency changes the plan.',
     fill=LIGHTAMBER, size=11.5, pad=0.16)
tf = tb(s, MARGIN, 6.10, CW, 0.4)
para(tf, True, 'Section 8 of the written outline lists the same four layers, with a licence link for each reused library.',
     size=10.5, color=MUTE, after=0)

# ========================================================= 8. contribution
s = slide('5 · Proposed features', 'Three contributions, and the rule that keeps them honest')
contrib = [
    ('Unified analytics', ['De-duplicates authorised logs and shows source, coverage and session detail — '
                           'but only where the records support it.',
                           'Where a session view is not supported, it shows a provider-level estimate instead of '
                           'inventing one.'], LIGHTBLUE, NAVY),
    ('Reproducible costs', ['Java applies dated official rates by provider, model, token class and service tier, '
                            'and keeps the price version behind every estimate.',
                            'New rates apply prospectively; old estimates keep the version they were computed with.'], LIGHTTEAL, TEAL),
    ('Evidence-based AI advice', ['Text or speech questions obtain calculated evidence from read-only tools; replies '
                                  'disclose sample size and missing data.',
                                  'It compares recorded cost, tokens and cache use — never answer quality, and it '
                                   'asks for no task ratings.'], LIGHTAMBER, AMBER),
]
cw, gap = (CW - 2 * 0.26) / 3, 0.26
for i, (head, body, fill, hc) in enumerate(contrib):
    card(s, MARGIN + i * (cw + gap), 1.70, cw, 2.34, head=head, body=body,
         fill=fill, head_color=hc, size=11, head_size=13.5, pad=0.16)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, 4.30, CW, 1.02,
      fill=LIGHTAMBER, line=AMBER, radius=0.06)
tf = s.shapes[-1].text_frame
tf.margin_left = tf.margin_right = Inches(0.20)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
para(tf, True, [('Scope rule:  ', True, AMBER),
                ('API-cost analytics and the home-screen assistant are core; forum and game are supporting '
                 'extensions. Live cross-agent status and system-wide overlays are outside the first release.',
                 False, INK)], size=12, after=0, line=1.14)
tf = tb(s, MARGIN, 5.54, CW, 0.75)
para(tf, True, [('Values are labelled.  ', True, NAVY),
                ('Every monetary figure is Estimated, Partial estimate or Unavailable — a cost that cannot be '
                 'computed is never shown as zero. Source-billed amounts are kept as separate reconciliation '
                 'evidence, not as replacements for the estimates.', False, INK)], size=11.5, after=0, line=1.16)

# ===================================================== 9. game (new feature)
s = slide('5 · Supporting feature A', 'A monthly tower-defence season built from tokens already logged')
s.shapes.add_picture(os.path.join(BASE, 'fig_game.png'), Inches(MARGIN), Inches(1.62), height=Inches(5.05))
l = 3.62
tf = tb(s, l, 1.62, PAGE_W - MARGIN - l, 5.05)
bullets = [
    ('It is a supporting tab.', 'Reached from the bottom bar, beside the forum. It is played in monthly seasons that match the billing cycle the provider subscriptions already follow.'),
    ('Tokens become resources.', 'Input, cache and output tokens already logged convert into separate resource pools at a fixed rate per million tokens.'),
    ('You plan the space yourself.', 'Towers and wall segments are placed and upgraded by hand on a build grid around the data-centre core, which expands in one direction.'),
    ('Enemies arrive from one side.', 'A wave enters from a single edge, so where each tower stands and whether the walls close decide how the wave goes.'),
    ('A wave starts when you start it.', 'Waves run in real time, so no offline simulation and no scheduled settlement is needed. The season resets each month.'),
]
first = True
for head, body in bullets:
    para(tf, first, [('▪  ', False, AMBER), (head + '  ', True, NAVY), (body, False, INK)],
         size=12.5, after=7, line=1.14)
    first = False
para(tf, False, [('Why it matters.  ', True, TEAL),
                 ('A day of agent work leaves a visible result without asking the user to spend differently. Playing '
                  'never changes usage, cost or budget totals, and re-importing the same logs grants no duplicate '
                  'resources — so the balance cannot be reached on invented data.',
                  False, INK)], size=12, after=0, before=6, line=1.16)

# ==================================================== 10. agent (new feature)
s = slide('5 · Core feature B', 'An assistant that answers from evidence, not opinion')
s.shapes.add_picture(os.path.join(BASE, 'fig_agent.png'), Inches(MARGIN), Inches(1.62), height=Inches(5.05))
tf = tb(s, l, 1.62, PAGE_W - MARGIN - l, 5.05)
bullets = [
    ('It lives on the home screen.', 'An "Ask AI / hold to talk" entry above the navigation opens an expandable panel on the analytics dashboard, where the figures it reasons about already are. It reads your logged usage and recent forum posts.'),
    ('Ask by text, or hold to talk.', 'Dictate a question, then review or edit the transcript before sending. The microphone is requested when used; denial or failure leaves typing available, and raw audio is not retained by default.'),
    ('A Java service does the work.', 'It calls one selected OpenAI, Xiaomi MiMo or DeepSeek pay-as-you-go model, which may request tools through function calling. Java executes and validates every call.'),
    ('Five read-only tools.', 'getUsageSummary · getBudgetStatus · compareAgentCosts · getForumHighlights · getMyThreads. The UID comes from the signed session token, never from a parameter.'),
    ('It admits what it does not know.', 'Replies state evidence, price versions and missing data; no logs means no cost-based answer. Its own tokens and cost sit in a separate ledger, excluded from coding-agent totals and game resources.'),
]
first = True
for head, body in bullets:
    para(tf, first, [('▪  ', False, AMBER), (head + '  ', True, NAVY), (body, False, INK)],
         size=12, after=6, line=1.13)
    first = False
para(tf, False, [('Boundary.  ', True, TEAL),
                 ('It can recommend but cannot run, route or interrupt a coding agent, and cannot enforce a hard '
                  'spending cap. It compares cost and token use, never answer quality.',
                  False, INK)], size=11.5, after=0, before=5, line=1.15)

# ============================================================ 11. data flow
s = slide('5 · How it is implemented', 'One path from logs to the app, and the four data aspects')
s.shapes.add_picture(os.path.join(BASE, 'fig_flow.png'), Inches((PAGE_W - 8.6) / 2), Inches(1.58), width=Inches(8.6))
aspects = [
    ('Data input', 'Profiles, billing mode and budgets; authorised logs, text or confirmed speech, posts and game actions.', NAVY, LIGHTBLUE),
    ('Data processing', 'Filter and de-duplicate logs; price usage; read-only summaries; rank posts and convert game resources.', TEAL, LIGHTTEAL),
    ('Data storage', 'Room holds local records. The server-side store holds user sync, budgets, game state and the assistant ledger; public posts are separate.', NAVY, LIGHTBLUE),
    ('Data output', 'Costs, heatmap, sessions, budgets and AI evidence; forum lists and real-time battle feedback.', AMBER, LIGHTAMBER),
]
cw, gap = (CW - 3 * 0.20) / 4, 0.20
for i, (head, body, hc, fill) in enumerate(aspects):
    card(s, MARGIN + i * (cw + gap), 3.34, cw, 1.56, head=head, body=body,
         fill=fill, head_color=hc, size=10.5, head_size=12.5, pad=0.14)
card(s, MARGIN, 5.04, CW, 1.02, head='Where it stops',
     body='Verify the log version and the billing mode; DSH may need JSONL or Zstandard decoding. The five tools '
          'enforce the caller\'s UID with bounded queries, and models receive summaries — never secrets or raw agent '
          'prompts. Missing logs mean no personalised cost claim, and there is no coding-agent control.',
     fill=LIGHTROW, head_color=NAVY, size=11.5, head_size=12.5)
tf = tb(s, MARGIN, 6.22, CW, 0.4)
para(tf, True, 'The same four aspects (input, processing, storage, output) run through the written outline, the '
               'implementation report, the code comments and the final demonstration.',
     size=10.5, color=MUTE, after=0)

# ============================================================ 12. proposed UI
s = slide('5 · Proposed interface', 'Home leads with analytics and the assistant entry')
s.shapes.add_picture(os.path.join(BASE, 'fig_ui.png'), Inches((PAGE_W - 11.1) / 2), Inches(1.52), width=Inches(11.1))
tf = tb(s, MARGIN, 6.26, CW, 0.4)
para(tf, True, 'Tap "Ask AI" to open the panel, or hold to dictate and confirm the transcript. Sessions expose '
               'source, date and token classes only where logs support them, and comparisons use matching filters. '
               'Forum and Game are supporting tabs, and assistant API costs are excluded from the home total.',
     size=10.5, color=MUTE, after=0)

# ============================================= 13. functional requirements
s = slide('5 · Proposed features', 'Functional requirements mapped to what already exists')
table(s, MARGIN, 1.60, [0.45, 2.30, 7.55, 1.593],
      [['Ser', 'Existing feature', 'TokenTrail improvement / addition', 'Type'],
       ['1', 'Cost dashboards',
        'Cross-provider dated estimates with source and coverage', 'Improved'],
       ['2', 'Session viewers',
        'Log-supported sessions and token-class breakdowns', 'Improved'],
       ['3', 'Billing budgets',
        'Cross-provider estimated budget and advisory alerts', 'Improved'],
       ['4', 'AI assistants',
        'Home AI: text and voice, usage evidence and a separate API ledger', 'New'],
       ['5', 'Forums',
        'Linked official news, posts and replies, and context for the advice', 'Improved'],
       ['6', 'Tower defence',
        'Monthly game resources from verified coding-agent API logs', 'New']],
      row_h=0.62, head_h=0.48, size=11)
tf = tb(s, MARGIN, 5.98, CW, 0.4)
para(tf, True, 'Six rows, in the same order as section 11 of the written outline. Only rows 4 and 6 are new; the '
               'rest improve a view the providers already offer one account at a time.',
     size=10.5, color=MUTE, after=0)

# =============================================================== 14. plan
s = slide('6 · Approach and work plan', 'Checkpoints from this week to the final submission')
table(s, MARGIN, 1.60, [1.42, 6.55, 3.923],
      [['Course weeks', 'Work', 'Checkpoint / acceptance'],
       ['1-3', 'Verify Codex/ZCode/DSH samples, billing modes, rate dimensions and UI; define access rules.',
        'Verified sources; isolated accounts.'],
       ['4-7', 'Build import, storage, dated pricing and analytics; define the five read-only tool contracts.',
        'Correct rates; idempotent import.'],
       ['8-9 · Alpha', 'Login → profile → import → cost → home.',
        'Separate real and sample data; source, GitHub link and MP4.'],
       ['10-12 · Beta', 'Wang: assistant and forum using Zhang\'s queries. Liu: settlement and the real-time game.',
        'Cited evidence; typing fallback; ranked posts; smooth waves.'],
       ['13-15 · Final', 'Validate advice, cost separation, game balance, accessibility and isolation; complete the documentation.',
        'Demo, tests, report and video; licence records.']],
      row_h=0.80, head_h=0.46, size=11)
tf = tb(s, MARGIN, 6.28, CW, 0.4)
para(tf, True, 'The team commits to GitHub throughout and adjusts scope based on tested progress.',
     size=10.5, color=MUTE, after=0)

# ============================================== 15. team, demo, validation
s = slide('6 · Evidence of completion', 'Who builds what, and how we will know it works')
team = [
    ('Zhang Li · 24107757', 'Authentication, log import and storage, dated pricing, budgets, the dashboard and the data-query interfaces.', LIGHTBLUE, NAVY),
    ('Wang Tingdong · 24107759', 'The home-screen AI assistant, its text/voice panel, the Java advice service, the five read-only tools and the assistant API ledger; forum news, posts, replies and ranking.', LIGHTTEAL, TEAL),
    ('Liu Zongrun · 24107745', 'The tower-defence engine and UI, resource settlement, monthly state, placement and upgrades, and enemy waves.', LIGHTAMBER, AMBER),
]
cw, gap = (CW - 2 * 0.24) / 3, 0.24
for i, (head, body, fill, hc) in enumerate(team):
    card(s, MARGIN + i * (cw + gap), 1.70, cw, 1.72, head=head, body=body,
         fill=fill, head_color=hc, size=11, head_size=12.5, pad=0.15)
card(s, MARGIN, 3.56, CW, 1.26, head='Demo path',
     body='Account isolation → tracked Codex, ZCode and DSH profiles → dated estimates from token logs → heatmap and '
          'budget warning → "Why did cost rise?" asked by text or confirmed speech, answered with evidence and its own '
          'assistant API cost → forum posts and replies → a player-started game wave.',
     fill=LIGHTROW, size=11.5, head_size=12.5)
card(s, MARGIN, 4.96, CW, 1.34, head='Validation',
     body='Test plan exclusion, log formats, token classes, dated and time-based rates, missing data, de-duplication '
          'and account/tool isolation. Replies are checked against evidence, text fallback and ledger separation, and '
          're-importing logs grants no duplicate game resources. Sample data never enters real totals.',
     fill=LIGHTROW, size=11.5, head_size=12.5)
tf = tb(s, MARGIN, 6.42, CW, 0.4)
para(tf, True, [('Repository: ', True, NAVY), (GH, False, TEAL, False, GH),
                ('   ·   UI/UX, accessibility and integration are shared across all three members.', False, MUTE)],
     size=11, after=0)

out = os.path.join(os.path.dirname(BASE), 'TokenTrail_Project_Outline_Slides.pptx')
_prs.save(out)
print('saved', out, 'slides:', len(_prs.slides._sldIdLst))
